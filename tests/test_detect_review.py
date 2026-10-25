import csv
import json
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

import cv2
import numpy as np
import pytest

from trafficwb.detect import Detection, _merge_across_classes, _nms_per_class, _touches_inner_edge, tile_grid
from trafficwb.export import IMAGE_COLUMNS, OBJECT_COLUMNS, build_tables, export, write_xlsx
from trafficwb.ingest import cycle
from trafficwb.models import ModelPack, ModelPackError, find, sha256_file
from trafficwb.process import Processor
from trafficwb.regions import RegionEditor
from trafficwb.review import Filter, ReviewService
from trafficwb.store import DataRoot, MissionConfig
from trafficwb.views import ViewService

W, H = 640, 360
OLD = datetime(2026, 10, 1, tzinfo=timezone.utc).timestamp()
RED, BLUE = (0, 0, 255), (255, 0, 0)


# ---------- synthetic data and a stand-in detector ----------

def scene(seed=1):
    rng = np.random.default_rng(seed)
    img = np.full((H, W, 3), 90, np.uint8)
    for _ in range(60):
        p1 = tuple(int(v) for v in rng.integers(0, [W, H]))
        p2 = tuple(int(v) for v in rng.integers(0, [W, H]))
        cv2.line(img, p1, p2, tuple(int(c) for c in rng.integers(0, 200, 3)), int(rng.integers(1, 4)))
    return img


def frame(k, cars=(), trucks=()):
    rng = np.random.default_rng(100 + k)
    img = cv2.add(scene(), rng.integers(0, 6, (H, W, 3), dtype=np.uint8))
    for (x, y) in cars:
        cv2.rectangle(img, (x, y), (x + 30, y + 20), RED, -1)
    for (x, y) in trucks:
        cv2.rectangle(img, (x, y), (x + 50, y + 30), BLUE, -1)
    return img


class FakeDetector:
    """Finds pure red (car) and pure blue (truck) rectangles. A black image raises, like a decoder failure."""

    def __init__(self):
        self.calls = 0
        self.fail_after = None

    def describe(self):
        return {"model": {"name": "fake", "version": "1", "sha256": "0" * 64},
                "preprocessing": {"note": "colour thresholds, no resize"}}

    def detect(self, img):
        self.calls += 1
        if self.fail_after is not None and self.calls > self.fail_after:
            raise KeyboardInterrupt  # simulates the process being killed mid-run
        if img.mean() < 3:
            raise RuntimeError("cannot decode")
        out = []
        for colour, cls in ((RED, "car"), (BLUE, "truck")):
            mask = np.all(np.abs(img.astype(int) - colour) <= 6, axis=2).astype(np.uint8)
            n, _, stats, _ = cv2.connectedComponentsWithStats(mask)
            for i in range(1, n):
                x, y, w, h, a = stats[i]
                if a > 100:
                    out.append(Detection((float(x), float(y), float(x + w), float(y + h)), cls, 0.9, {cls: 0.9}))
        return out


def ts(i):
    return (datetime(2026, 10, 21, tzinfo=timezone.utc) + timedelta(minutes=10 * i)).strftime("%Y%m%dT%H%M%SZ")


def write(src, name, img):
    p = src / name
    cv2.imwrite(str(p), img, [cv2.IMWRITE_JPEG_QUALITY, 98])
    os.utime(p, (OLD, OLD))


@pytest.fixture
def mission(tmp_path):
    src, out = tmp_path / "src", tmp_path / "out"
    src.mkdir()
    out.mkdir()
    root = DataRoot(tmp_path / "data")
    m = root.create_mission(MissionConfig(
        name="Tie 3 – בדיקה", input_type="images", mode="general_traffic", country="IL", iana_zone="Asia/Jerusalem",
        source_dir=str(src), include_subdirs=False, processing="continuous", output_dir=str(out),
        export_formats=["csv", "xlsx"]), "test")
    return m, src, out


def vid_of(m, name):
    return m.con.execute("SELECT v.id FROM source_version v JOIN source_file f ON f.id=v.source_file_id "
                         "WHERE f.rel_path=?", (name,)).fetchone()[0]


def populate(src, n=4):
    for k in range(n):
        write(src, f"0045A_{ts(k)}.jpg", frame(k, cars=[(100, 250), (400, 100)], trucks=[(250, 200)]))


# ---------- detector helpers ----------

def test_tiles_cover_frame_with_overlap():
    tiles = tile_grid(1280, 720, 3, 2, 0.2)
    assert len(tiles) == 6
    assert min(t[0] for t in tiles) == 0 and max(t[2] for t in tiles) == 1280
    assert min(t[1] for t in tiles) == 0 and max(t[3] for t in tiles) == 720
    assert tiles[0][2] > tiles[1][0]  # neighbours overlap


def test_tile_edge_boxes_dropped_except_at_frame_edges():
    tile = (400, 0, 900, 400)
    assert _touches_inner_edge((401, 50, 450, 90), tile, (1280, 720))      # cut by left inner edge
    assert not _touches_inner_edge((450, 0, 500, 40), tile, (1280, 720))   # touches the frame top, kept
    assert not _touches_inner_edge((500, 100, 560, 140), tile, (1280, 720))


def test_nms_merges_duplicates_per_class():
    a = Detection((10, 10, 50, 40), "car", 0.9, {})
    b = Detection((12, 11, 51, 41), "car", 0.6, {})
    c = Detection((12, 11, 51, 41), "truck", 0.5, {})
    kept = _nms_per_class([a, b, c], 0.45)
    assert a in kept and b not in kept and c in kept
    assert _merge_across_classes(kept, 0.65) == [a]   # same vehicle as car and truck counts once


# ---------- model packs ----------

def _pack(tmp_path, data=b"weights", sha=None):
    d = tmp_path / "models" / "m" / "1.0"
    d.mkdir(parents=True)
    (d / "model.onnx").write_bytes(data)
    card = {"name": "m", "version": "1.0", "task": "vehicle_detection", "license": "Apache-2.0",
            "sha256": sha or sha256_file(d / "model.onnx")}
    (d / "model.json").write_text(json.dumps(card))
    return d


def test_model_pack_checksum_verified(tmp_path):
    d = _pack(tmp_path)
    assert ModelPack.load(d).name == "m"
    assert find("m", root=tmp_path / "models").version == "1.0"
    (d / "model.onnx").write_bytes(b"tampered")
    with pytest.raises(ModelPackError, match="do not match"):
        ModelPack.load(d)


def test_model_pack_missing_parts(tmp_path):
    with pytest.raises(ModelPackError):
        find("absent", root=tmp_path)
    d = _pack(tmp_path)
    (d / "model.onnx").unlink()
    with pytest.raises(ModelPackError, match="missing"):
        ModelPack.load(d)


# ---------- processing ----------

def test_processing_stores_immutable_detections_and_run_record(mission):
    m, src, _ = mission
    populate(src)
    _, run_id = cycle(m)
    rep = Processor(m, FakeDetector()).process_run(run_id)
    assert (rep.done, rep.failed, rep.detections) == (4, 0, 12)
    run = m.con.execute("SELECT * FROM run WHERE id=?", (run_id,)).fetchone()
    assert run["state"] == "completed"
    assert json.loads(run["model_versions"])["vehicle_detection"]["name"] == "fake"
    assert json.loads(run["preprocessing"])["region_rule"]["rule"] == "bottom_center"
    row = m.con.execute("SELECT typeof(x1), typeof(score), cls FROM detection LIMIT 1").fetchone()
    assert row[0] == "real" and row[1] == "real"
    assert m.con.execute("SELECT COUNT(*) FROM model_version").fetchone()[0] == 1
    # Processing again does nothing: no duplicates.
    assert Processor(m, FakeDetector()).process_all() == []
    assert m.con.execute("SELECT COUNT(*) FROM detection").fetchone()[0] == 12


def test_failed_image_recorded_and_run_continues(mission):
    m, src, _ = mission
    populate(src, 3)
    write(src, f"0045A_{ts(9)}.jpg", np.zeros((H, W, 3), np.uint8))
    _, run_id = cycle(m)
    rep = Processor(m, FakeDetector()).process_run(run_id)
    assert rep.done == 3 and rep.failed == 1
    bad = m.con.execute("SELECT * FROM run_item WHERE processing_state='failed'").fetchone()
    assert bad["error_code"] == "RuntimeError"
    assert m.con.execute("SELECT state FROM run").fetchone()[0] == "completed_with_errors"
    assert m.con.execute("SELECT COUNT(*) FROM notification WHERE code='processing_failed'").fetchone()[0] == 1


def test_interrupted_run_resumes_without_duplicates(mission):
    m, src, _ = mission
    populate(src, 4)
    _, run_id = cycle(m)
    det = FakeDetector()
    det.fail_after = 2
    with pytest.raises(KeyboardInterrupt):
        Processor(m, det).process_run(run_id)
    assert m.con.execute("SELECT COUNT(*) FROM run_item WHERE processing_state='done'").fetchone()[0] == 2
    reps = Processor(m, FakeDetector()).process_all()
    assert reps[0].done == 2
    assert m.con.execute("SELECT COUNT(*) FROM detection").fetchone()[0] == 12
    assert m.con.execute("SELECT COUNT(DISTINCT run_item_id) FROM detection").fetchone()[0] == 4


def test_exclusion_region_applied_and_counts(mission):
    m, src, _ = mission
    populate(src, 4)
    cycle(m)
    vs = ViewService(m)
    vs.assign_or_propose()
    view = vs.views()[0]["id"]
    # Exclude the top-right car (400,100)-(430,120); bottom-centre is (415,120).
    RegionEditor(m.con).save(view, "exclude", [(380, 80), (460, 80), (460, 140), (380, 140)], "lior")
    Processor(m, FakeDetector()).process_all()
    rows = ReviewService(m).results()
    assert [r.vehicle_count for r in rows] == [2, 2, 2, 2]
    assert m.con.execute("SELECT COUNT(*) FROM detection WHERE in_region=0").fetchone()[0] == 4


# ---------- review ----------

def _processed(mission, n=4):
    m, src, out = mission
    populate(src, n)
    cycle(m)
    Processor(m, FakeDetector()).process_all()
    return m, src, out, ReviewService(m)


def test_corrections_are_revisions_and_machine_output_unchanged(mission):
    m, src, _, rs = _processed(mission)
    v = vid_of(m, f"0045A_{ts(0)}.jpg")
    before = [tuple(r) for r in m.con.execute("SELECT * FROM detection ORDER BY id")]
    objs = rs.objects(v)
    car = next(o for o in objs if o.cls == "car")
    truck = next(o for o in objs if o.cls == "truck")
    rs.remove(v, car.uid, "lior", note="shadow, not a car")
    rs.reclass(v, truck.uid, "bus", "lior")
    added = rs.add(v, (500, 300, 540, 330), "motorcycle", "lior")
    now = {o.uid: o for o in rs.objects(v)}
    assert now[car.uid].status == "removed" and now[car.uid].correction == "removed"
    assert now[truck.uid].cls == "bus" and now[truck.uid].machine_cls == "truck" and now[truck.uid].correction == "reclassed"
    assert now[added].origin == "human" and now[added].correction == "added"
    assert sum(o.counted for o in now.values()) == 3
    rs.restore(v, car.uid, "lior")
    assert sum(o.counted for o in rs.objects(v)) == 4
    assert [tuple(r) for r in m.con.execute("SELECT * FROM detection ORDER BY id")] == before
    assert m.con.execute("SELECT COUNT(*) FROM annotation_revision").fetchone()[0] == 4
    with pytest.raises(ValueError):
        rs.add(v, (600, 300, 700, 340), "car", "lior")   # beyond the 640-pixel width
    with pytest.raises(ValueError):
        rs.reclass(v, truck.uid, "spaceship", "lior")


def test_moving_box_recomputes_region_membership(mission):
    m, src, _ = mission
    populate(src, 4)
    cycle(m)
    vs = ViewService(m)
    vs.assign_or_propose()
    view = vs.views()[0]["id"]
    RegionEditor(m.con).save(view, "exclude", [(0, 0), (200, 0), (200, 100), (0, 100)], "lior")
    Processor(m, FakeDetector()).process_all()
    rs = ReviewService(m)
    v = vid_of(m, f"0045A_{ts(1)}.jpg")
    car = next(o for o in rs.objects(v) if o.cls == "car" and o.box[1] > 200)
    assert car.in_region == 1
    rs.move(v, car.uid, (50, 40, 80, 60), "lior")
    moved = next(o for o in rs.objects(v) if o.uid == car.uid)
    assert moved.in_region == 0 and not moved.counted and moved.correction == "moved"


def test_approve_and_edit_reopens(mission):
    m, src, _, rs = _processed(mission)
    v = vid_of(m, f"0045A_{ts(0)}.jpg")
    rs.approve(v, "lior")
    assert rs.review_status(v) == "approved"
    rs.remove(v, rs.objects(v)[0].uid, "lior")
    assert rs.review_status(v) == "unreviewed"
    assert [r[0] for r in m.con.execute("SELECT decision FROM review_decision ORDER BY id")] == ["approved", "reopened"]


def test_approve_all_is_bulk_fixed_scope_and_excludes_later_data(mission):
    m, src, _, rs = _processed(mission)
    v0 = vid_of(m, f"0045A_{ts(0)}.jpg")
    rs.approve(v0, "lior")
    prev = rs.approve_all_preview()
    assert prev["count"] == 3                       # already-approved image is not re-counted
    # New data arrives after the operator saw the count.
    write(src, f"0045A_{ts(20)}.jpg", frame(20, cars=[(100, 250)]))
    cycle(m)
    Processor(m, FakeDetector()).process_all()
    bid = rs.approve_all_commit(prev, "lior")
    statuses = {r.filename: r.review_status for r in rs.results()}
    assert statuses[f"0045A_{ts(20)}.jpg"] == "unreviewed"
    assert statuses[f"0045A_{ts(0)}.jpg"] == "approved"
    assert list(statuses.values()).count("bulk_approved_unreviewed") == 3
    assert m.con.execute("SELECT item_count FROM bulk_approval WHERE id=?", (bid,)).fetchone()[0] == 3


def test_unprocessed_image_cannot_be_approved(mission):
    m, src, _ = mission
    populate(src, 1)
    cycle(m)
    rs = ReviewService(m)
    with pytest.raises(ValueError):
        rs.approve(vid_of(m, f"0045A_{ts(0)}.jpg"), "lior")
    assert rs.results()[0].processing_state == "queued"


def test_results_filter_sort_and_position(mission):
    m, src, _, rs = _processed(mission)
    rs.add(vid_of(m, f"0045A_{ts(2)}.jpg"), (500, 300, 540, 330), "car", "lior")
    rows = rs.results(sort="vehicle_count", descending=True)
    assert rows[0].filename == f"0045A_{ts(2)}.jpg" and rows[0].vehicle_count == 4
    assert rows[0].capture_local.endswith("+03:00")   # Israel summer time on 21 Oct 2026
    assert len(rs.results(Filter(search=ts(1)))) == 1
    assert len(rs.results(Filter(min_vehicles=4))) == 1
    assert len(rs.results(limit=2, offset=1)) == 2
    pos = rs.position(vid_of(m, f"0045A_{ts(1)}.jpg"))
    assert pos["index"] == 2 and pos["total"] == 4
    assert pos["prev"] == vid_of(m, f"0045A_{ts(0)}.jpg") and pos["next"] == vid_of(m, f"0045A_{ts(2)}.jpg")


# ---------- export ----------

def test_export_csv_and_xlsx_traceable_without_overwrite(mission):
    from openpyxl import load_workbook

    m, src, out, rs = _processed(mission)
    v = vid_of(m, f"0045A_{ts(0)}.jpg")
    car = next(o for o in rs.objects(v) if o.cls == "car")
    rs.remove(v, car.uid, "lior")
    rs.approve(v, "lior")
    first = export(m, "lior")
    second = export(m, "lior")
    assert len(first) == 3 and len(second) == 3
    assert not set(first) & set(second)            # never overwritten
    imgs = next(p for p in first if p.name.endswith("_images.csv"))
    with open(imgs, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == IMAGE_COLUMNS
    assert rows[0]["file"].startswith("0045A_")    # leading zeros kept
    assert rows[0]["utc_offset"] == "+03:00" and rows[0]["time_zone"] == "Asia/Jerusalem"
    assert rows[0]["review_status"] == "approved" and rows[0]["vehicles_in_image"] == "2"
    objs = next(p for p in first if p.name.endswith("_objects.csv"))
    with open(objs, encoding="utf-8-sig", newline="") as f:
        orows = list(csv.DictReader(f))
    removed = [r for r in orows if r["final_status"] == "removed"]
    assert len(removed) == 1 and removed[0]["machine_class"] == "car" and removed[0]["counted"] == "no"
    assert {r["origin"] for r in orows} == {"machine"}
    assert "– בדיקה" in imgs.name or "בדיקה" in imgs.name or imgs.name.startswith("Tie_3")
    xl = next(p for p in first if p.suffix == ".xlsx")
    wb = load_workbook(xl, read_only=True)
    ws = wb["images"]
    vals = list(ws.iter_rows(values_only=True))
    assert vals[0] == tuple(IMAGE_COLUMNS) and vals[1][3].startswith("0045A_")
    assert m.con.execute("SELECT COUNT(*) FROM export WHERE status='completed'").fetchone()[0] == 4
    assert not list(out.glob("*.partial"))


def test_xlsx_never_truncates_rows(tmp_path):
    images = [[str(i)] + [""] * (len(IMAGE_COLUMNS) - 1) for i in range(25)]
    p = write_xlsx(tmp_path, "t", images, [], max_rows=11)[0]
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    names = [n for n in wb.sheetnames if n.startswith("images")]
    assert names == ["images", "images_2", "images_3"]
    total = sum(len(list(wb[n].iter_rows(values_only=True))) - 1 for n in names)
    assert total == 25


def test_export_failure_is_recorded(mission, tmp_path):
    m, src, out, rs = _processed(mission, 3)
    blocker = tmp_path / "not_a_dir"
    blocker.write_text("x")
    with pytest.raises(Exception):
        export(m, "lior", out_dir=blocker)
    assert m.con.execute("SELECT status FROM export ORDER BY id DESC").fetchone()[0] == "failed"
    assert m.con.execute("SELECT COUNT(*) FROM notification WHERE code='export_failed'").fetchone()[0] == 1


# ---------- the real detector (needs the model pack) ----------

def _real_pack():
    try:
        return find("yolox-s")
    except ModelPackError:
        return None


@pytest.mark.skipif(_real_pack() is None, reason="model pack yolox-s not installed (run scripts/fetch_models.py)")
def test_real_detector_runs_and_records_preprocessing():
    from trafficwb.detect import YoloxDetector
    det = YoloxDetector(_real_pack())
    assert det.detect(np.full((720, 1280, 3), 120, np.uint8)) == []
    d = det.describe()
    assert d["model"]["name"] == "yolox-s" and d["preprocessing"]["tiles"]["tiles"] == (3, 2)
    sample = os.environ.get("TWB_SAMPLE_DIR")
    if sample:
        from trafficwb.views import read_image
        found = det.detect(read_image(Path(sample) / "C0450701_20261021T071528Z.jpg"))
        assert sum(d.cls == "car" for d in found) >= 8
        assert all(isinstance(v, float) for d in found for v in d.box)
