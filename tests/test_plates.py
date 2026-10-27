"""Plates: normalisation, layouts, readability, the plate stage, corrections, evaluation gate and export (IMG-04/05)."""
import csv
import os
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient

from trafficwb.detect import Detection
from trafficwb.models import ModelPackError, find
from trafficwb.plate_layout import ARABIC_DIGITS, PERSIAN_DIGITS, ReadRules, assess, cer, layout_of, normalize
from trafficwb.plates import OcrOut, PlateBox
from trafficwb.web.app import create_app

OLD = datetime(2026, 10, 1, tzinfo=timezone.utc).timestamp()
H1 = {"X-TWB": "1"}
PAD = "_"


def slots(text, conf=0.99, n=10, low=()):
    out = [(c, 0.3 if i in low else conf) for i, c in enumerate(text)]
    return out + [(PAD, 0.99)] * (n - len(out))


# ------------------------------------------------------------------ text rules

def test_normalize_keeps_order_zeros_and_script():
    assert normalize("12-345-67") == "1234567"
    assert normalize("0012 345") == "0012345"
    assert normalize("ab-123") == "AB123"
    assert normalize(PERSIAN_DIGITS[1] + PERSIAN_DIGITS[2] + "ب" + PERSIAN_DIGITS[3] * 3) == "12ب333"
    assert normalize(ARABIC_DIGITS[0] + ARABIC_DIGITS[7]) == "07"


def test_layouts_and_display():
    r = assess(slots("1234567"), PAD, "IL")
    assert (r.normalized, r.display, r.layout, r.readability) == ("1234567", "12-345-67", "private-7", "full")
    r = assess(slots("01234567"), PAD, "IL")
    assert (r.display, r.layout) == ("012-34-567", "private-8")
    assert layout_of("AB1234", "LV").name == "standard"
    assert layout_of("RIGA2026", "LV").name == "personalised"
    assert layout_of("12ب34567", "IR").name == "private"
    assert layout_of("12ت34567", "IR").name == "other"          # taxi letter
    assert layout_of("A123B", "LV").name == "trailer"
    assert layout_of("AAAA12", "LV") is None                      # no symbol four times in a row
    assert layout_of("A1234", "LV").name == "trailer"             # one letter is a trailer, not a car
    assert layout_of("B123456", "LB").name == "private" and layout_of("RL123456", "LB").name == "rl"
    assert layout_of("C20154", "LB").name == "consular"
    assert layout_of("D20154", "LB").name == "diplomatic"
    assert layout_of("P123456", "LB").name == "public-transport"
    assert layout_of("M123456", "LB").name == "commercial-service-family"
    assert layout_of("MP45", "LB").name == "parliament" and layout_of("129", "LB") is None
    assert assess(slots("123456"), PAD, "IL").display == "123-456"
    assert normalize("۱۲هـ۳۴۵۶۷") == "12ه34567"                   # tatweel removed, letter kept


def test_unsure_characters_are_never_filled():
    r = assess(slots("1234567", low=(3,)), PAD, "IL")
    assert r.normalized == "123?567" and r.readability == "partial" and r.layout is None
    assert r.raw_text == "1234567"                     # the reader's output is kept as emitted
    # An unsure pad inside the text is an unsure character, not the end of the plate.
    s = slots("12") + []
    s = [("1", .99), ("2", .99), (PAD, .4), ("4", .99), (PAD, .99)] + [(PAD, .99)] * 5
    r = assess(s, PAD, "LV")
    assert r.normalized == "12?4" and r.readability == "partial"
    # Nothing sure at all: unreadable, with no text.
    r = assess(slots("1234567", conf=0.2), PAD, "IL")
    assert r.readability == "unreadable" and r.normalized == ""
    # Too short for the country: partial, with the reason.
    r = assess(slots("123"), PAD, "IL")
    assert r.readability == "partial" and "length" in r.reasons[0]


def test_cer():
    assert cer("1234567", "1234567") == 0
    assert cer("1234567", "123?567") == pytest.approx(1 / 7)
    assert cer("12-345-67", "1234567") == 0


# ------------------------------------------------------------------ stage and API

W, H = 960, 540
# vehicle box, plate box (x, y, w, h), plate text for the stand-in reader
VEHICLES = [
    ((100, 200, 300, 360), (140, 300, 120, 26), slots("0123456")),          # full, leading zero
    ((400, 200, 620, 380), (440, 320, 140, 30), slots("7654321", low=(2,))),  # one unsure character
    ((700, 200, 860, 340), (760, 280, 40, 40), slots("T1")),                 # square box: not a plate shape
    ((880, 480, 900, 495), None, None),                                      # vehicle too small
]


def picture(k):
    img = np.full((H, W, 3), 80, np.uint8)
    img[:] += np.random.default_rng(k).integers(0, 8, (H, W, 3), dtype=np.uint8)
    for vb, pb, _ in VEHICLES:
        cv2.rectangle(img, vb[:2], vb[2:], (60, 40, 150), -1)
        if pb:
            x, y, w, h = pb
            cv2.rectangle(img, (x, y), (x + w, y + h), (0, 220, 255), -1)
    return img


class FixedDetector:
    def describe(self):
        return {"model": {"name": "fixed", "version": "1", "sha256": "0" * 64}, "preprocessing": {}}

    def detect(self, img):
        return [Detection(tuple(float(v) for v in vb), "car", 0.9, {"car": 0.9}) for vb, _, _ in VEHICLES]


class YellowFinder:
    """Finds the yellow plate rectangles."""

    def describe(self):
        return {"model": {"name": "yellow-finder", "version": "1", "sha256": "1" * 64}, "preprocessing": {}}

    def find(self, crop):
        m = np.all(np.abs(crop.astype(int) - (0, 220, 255)) <= 50, axis=2).astype(np.uint8)
        n, _, st, _ = cv2.connectedComponentsWithStats(m)
        return [PlateBox((float(x), float(y), float(x + w), float(y + h)), 0.9) for x, y, w, h, a in st[1:] if a > 20]


class SizeReader:
    """Returns the text planned for a plate of this size."""

    def __init__(self):
        self.calls = 0

    def describe(self):
        return {"model": {"name": "size-reader", "version": "1", "sha256": "2" * 64}, "preprocessing": {}}

    def read(self, plate):
        self.calls += 1
        h, w = plate.shape[:2]
        for _, pb, s in VEHICLES:
            if pb and abs(pb[2] - w) <= 3 and abs(pb[3] - h) <= 3:
                return OcrOut(s, PAD, "Israel", 0.9)
        return OcrOut(slots(""), PAD, None, None)


def ts(i):
    return (datetime(2026, 10, 21, tzinfo=timezone.utc) + timedelta(minutes=10 * i)).strftime("%Y%m%dT%H%M%SZ")


@pytest.fixture
def api(tmp_path):
    src, out = tmp_path / "src", tmp_path / "out"
    src.mkdir()
    out.mkdir()
    for k in range(2):
        p = src / f"CAM5_{ts(k)}.png"
        cv2.imwrite(str(p), picture(k))
        os.utime(p, (OLD, OLD))
    reader = SizeReader()
    app = create_app(tmp_path / "data", detector_factory=lambda: (FixedDetector(), {"task": "vehicle_detection"}),
                     plate_factory=lambda: (YellowFinder(), reader), start_monitoring=False, stability_seconds=0)
    c = TestClient(app)
    def make(country="IL", mode="plate_recognition"):
        body = {"name": f"P {country}", "mode": mode, "country": country, "iana_zone": "Asia/Jerusalem",
                "source_dir": str(src), "include_subdirs": False, "processing": "one_time", "output_dir": str(out),
                "export_formats": ["csv", "xlsx"]}
        r = c.post("/api/missions", json=body, headers=H1)
        assert r.status_code == 200, r.text
        mid = r.json()["id"]
        c.post(f"/api/missions/{mid}/job", headers=H1)
        wait(c, mid)
        return mid
    return c, make, out, reader


def wait(c, mid, timeout=60):
    t = time.time()
    while time.time() - t < timeout:
        j = c.get(f"/api/missions/{mid}").json()["job"]
        if j and j["state"] in ("finished", "cancelled", "failed"):
            assert j["state"] == "finished", j
            return j
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def by_box(plates):
    return {tuple(p["vehicle_box"]): p for p in plates}


def test_plate_stage_outcomes(api):
    c, make, _, reader = api
    mid = make()
    r = c.get(f"/api/missions/{mid}/plates").json()
    assert r["total"] == 8 and r["country"]["enabled"] is False and r["country"]["reader_covers"] is True
    first = [p for p in r["plates"] if p["source_version_id"] == r["plates"][0]["source_version_id"]]
    v = by_box(first)
    a = v[(100.0, 200.0, 300.0, 360.0)]
    assert a["outcome"] == "read" and a["machine_text"] == "0123456" and a["machine_display"] == "01-234-56"
    assert a["machine_readability"] == "full" and a["evaluation_only"] is True and a["layout"] == "private-7"
    assert a["box"] == pytest.approx([140, 300, 260, 326], abs=1.5)
    b = v[(400.0, 200.0, 620.0, 380.0)]
    assert b["machine_text"] == "76?4321" and b["machine_readability"] == "partial" and b["raw_text"] == "7654321"
    sq = v[(700.0, 200.0, 860.0, 340.0)]
    assert sq["machine_readability"] == "partial" and any("plate-shaped" in x for x in sq["reasons"])
    assert v[(880.0, 480.0, 900.0, 495.0)]["outcome"] == "vehicle_too_small"
    assert c.get(f"/api/missions/{mid}/plates/{a['id']}/crop").headers["content-type"] == "image/png"
    assert c.get(f"/api/missions/{mid}/plates?readability=partial").json()["total"] == 4
    assert c.get(f"/api/missions/{mid}/plates?search=01-23").json()["total"] == 2
    # A second job reads nothing again.
    calls = reader.calls
    c.post(f"/api/missions/{mid}/job", headers=H1)
    wait(c, mid)
    assert reader.calls == calls and c.get(f"/api/missions/{mid}/plates").json()["total"] == 8


def test_country_without_reader_finds_but_does_not_read(api):
    c, make, _, reader = api
    mid = make("IR")
    ps = c.get(f"/api/missions/{mid}/plates").json()
    a = by_box(ps["plates"])[(100.0, 200.0, 300.0, 360.0)]
    assert a["outcome"] == "read" and a["box"] is not None and a["raw_text"] is None
    assert a["machine_readability"] == "unreadable" and "no reader" in a["reasons"][0]
    assert ps["country"]["reader_covers"] is False and reader.calls == 0


def test_corrections_are_append_only_and_never_guess(api):
    c, make, out, _ = api
    mid = make()
    v = by_box([p for p in c.get(f"/api/missions/{mid}/plates").json()["plates"]][:4])
    a, b = v[(100.0, 200.0, 300.0, 360.0)], v[(400.0, 200.0, 620.0, 380.0)]
    rv = lambda pid, **body: c.post(f"/api/missions/{mid}/plates/{pid}/revise", json=body, headers=H1)
    assert rv(a["id"], action="confirm").json()["review"] == "confirmed"
    assert rv(b["id"], action="confirm").status_code == 400                     # partial cannot be confirmed as is
    assert rv(b["id"], action="correct", text="76?4321").status_code == 400     # no guessing in corrections either
    persian = PERSIAN_DIGITS[7] + PERSIAN_DIGITS[6] + "54321"
    got = rv(b["id"], action="correct", text=persian).json()
    assert got["text"] == "7654321" and got["review"] == "corrected" and got["machine_text"] == "76?4321"
    sq = v[(700.0, 200.0, 860.0, 340.0)]
    assert rv(sq["id"], action="unreadable", note="road sign").json()["readability"] == "unreadable"
    assert rv(sq["id"], action="reopen").json()["review"] == "pending"
    con_rows = c.get(f"/api/missions/{mid}/plates?review=corrected").json()["total"]
    assert con_rows == 1
    files = c.post(f"/api/missions/{mid}/plates/export", json={}, headers=H1).json()["files"]
    csvf = next(f for f in files if f.endswith(".csv"))
    rows = list(csv.DictReader(open(csvf, encoding="utf-8-sig")))
    row = next(r for r in rows if r["machine_text"] == "0123456")
    assert row["plate_text"] == "0123456" and row["evaluation_only"] == "yes" and row["review"] == "confirmed"
    from openpyxl import load_workbook
    ws = load_workbook(next(f for f in files if f.endswith(".xlsx")))["plates"]
    hdr = [x.value for x in ws[1]]
    vals = [r[hdr.index("machine_text")].value for r in ws.iter_rows(min_row=2)]
    assert "0123456" in vals                                                     # leading zero kept as text


def test_evaluation_gate(api):
    c, make, _, _ = api
    mid = make()
    assert c.post(f"/api/missions/{mid}/plates/country/enable", json={"enabled": True}, headers=H1).status_code == 400
    ps = c.get(f"/api/missions/{mid}/plates").json()["plates"]
    for p in ps:
        if p["machine_readability"] == "full":
            c.post(f"/api/missions/{mid}/plates/{p['id']}/revise", json={"action": "confirm"}, headers=H1)
    ev = c.post(f"/api/missions/{mid}/plates/country/evaluate", headers=H1).json()
    assert ev["reviewed_plates"] == 2 and ev["exact_match"] == 1.0 and ev["passed"] is False   # fewer than 100 plates
    assert c.post(f"/api/missions/{mid}/plates/country/enable", json={"enabled": True}, headers=H1).status_code == 400


def test_evaluation_counts_guesses(tmp_path, api):
    c, make, _, _ = api
    mid = make()
    from trafficwb.plates import PlateService
    from trafficwb.store import DataRoot
    m = DataRoot(tmp_path / "data").open_mission(mid)
    svc = PlateService(m)
    full = [p for p in svc.plates() if p.machine_readability == "full"]
    svc.revise(full[0].id, "confirm", "lior")
    svc.revise(full[1].id, "unreadable", "lior")         # the machine read a plate the operator cannot read
    ev = svc.evaluate("lior", min_plates=1, min_exact=0.5, max_cer=0.5)
    assert ev["machine_guesses_on_unreadable"] == 1 and ev["passed"] is False
    svc.revise(full[1].id, "confirm", "lior")
    ev = svc.evaluate("lior", min_plates=1, min_exact=0.5, max_cer=0.5)
    assert ev["passed"] is True
    st = svc.set_enabled(True, "lior")
    assert st["enabled"] is True and all(not p.evaluation_only for p in svc.plates())


def test_plates_only_in_plate_missions(api):
    c, make, _, _ = api
    mid = make(mode="general_traffic")
    assert c.get(f"/api/missions/{mid}/plates").status_code == 400


# ------------------------------------------------------------------ installed reader

def _plate_image(text, bg=(0, 200, 255)):
    img = np.full((110, 520, 3), bg, np.uint8)
    cv2.rectangle(img, (0, 0), (519, 109), (0, 0, 0), 3)
    cv2.putText(img, text, (20, 85), cv2.FONT_HERSHEY_SIMPLEX, 2.6, (0, 0, 0), 7, cv2.LINE_AA)
    return img


def test_installed_reader_on_rendered_plates():
    try:
        from trafficwb.plates import SlotPlateReader
        reader = SlotPlateReader(find("plate-ocr-cct-s-v2-global"))
    except ModelPackError:
        pytest.skip("plate reader pack not installed")
    for text, country, want in [("12-345-67", "IL", "1234567"), ("123-45-678", "IL", "12345678"), ("AB-1234", "LV", "AB1234")]:
        o = reader.read(_plate_image(text, (255, 255, 255) if country == "LV" else (0, 200, 255)))
        r = assess(o.slots, o.pad, country, ReadRules())
        assert r.normalized == want and r.readability == "full", (text, r)
    # A blank plate gives no confident text.
    o = reader.read(np.full((110, 520, 3), 200, np.uint8))
    assert assess(o.slots, o.pad, "IL").readability != "full"


def test_operator_plate_box_where_the_finder_missed(api):
    c, make, out, _ = api
    mid = make()
    v = by_box([p for p in c.get(f"/api/missions/{mid}/plates").json()["plates"]][:4])
    small = v[(880.0, 480.0, 900.0, 495.0)]                     # the finder never looked: vehicle too small
    a = v[(100.0, 200.0, 300.0, 360.0)]
    rv = lambda pid, **body: c.post(f"/api/missions/{mid}/plates/{pid}/revise", json=body, headers=H1)
    assert "not on this vehicle" in rv(small["id"], action="correct", text="12345", box=[100, 300, 140, 320]).json()["message"]
    assert "outside the image" in rv(small["id"], action="correct", text="12345", box=[885, 485, 970, 494]).json()["message"]
    assert rv(small["id"], action="correct", text="12345", box=[884, 488, 886, 489]).status_code == 400   # too small
    got = rv(small["id"], action="correct", text="12345", box=[882, 486, 898, 493]).json()
    assert got["operator_box"] == [882, 486, 898, 493] and got["text"] == "12345" and got["outcome"] == "vehicle_too_small"
    # The operator can also move a box the finder got wrong.
    got = rv(a["id"], action="correct", text="0123456", box=[141, 301, 259, 325]).json()
    assert got["operator_box"] == [141, 301, 259, 325] and got["box"] == pytest.approx([140, 300, 260, 326], abs=1.5)
    assert rv(a["id"], action="reopen").json()["operator_box"] is None
    ev = c.post(f"/api/missions/{mid}/plates/country/evaluate", headers=H1).json()
    assert ev["finder_missed"] == 1
    files = c.post(f"/api/missions/{mid}/plates/export", json={"formats": ["csv"]}, headers=H1).json()["files"]
    rows = list(csv.DictReader(open(files[0], encoding="utf-8-sig")))
    row = next(r for r in rows if r["plate_text"] == "12345")
    assert (row["operator_x1"], row["operator_x2"], row["plate_x1"]) == ("882", "898", "")
